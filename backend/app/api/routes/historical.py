"""Historical fare snapshot import and summary endpoints.

Imports the checked-in Kaggle 2019 workbook once, preserving it as a historical
reference. It is intentionally excluded from the live price index.
"""
from datetime import datetime, timezone
from pathlib import Path
import hashlib
from datetime import date

from openpyxl import load_workbook
from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.auth import require_admin, get_current_user
from app.core.database import get_db
from app.models.fare import FareObservation

router = APIRouter(prefix="/historical", tags=["historical"])

CITY_CODES = {
    "Delhi": "DEL", "New Delhi": "DEL", "Mumbai": "BOM", "Banglore": "BLR",
    "Bangalore": "BLR", "Kolkata": "CCU", "Chennai": "MAA", "Cochin": "COK",
    "Hyderabad": "HYD",
}
WORKBOOK = Path(__file__).resolve().parents[4] / "src" / "imports" / "Data_Train.xlsx"
SOURCE = "kaggle-flight-fare-prediction-mh"
ORIGIN = "HISTORICAL_PUBLIC"


def _date(value) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%Y %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return "2019-01-01"


@router.get("/summary")
async def historical_summary(
    db: AsyncSession = Depends(get_db),
    _: dict = Depends(get_current_user),
):
    q = select(FareObservation).where(FareObservation.data_origin == ORIGIN)
    total = await db.scalar(select(func.count()).select_from(q.subquery())) or 0
    if not total:
        return {"status": "NOT_IMPORTED", "records": 0, "source": SOURCE, "data_origin": ORIGIN}
    stats = (await db.execute(select(
        func.min(FareObservation.travel_date), func.max(FareObservation.travel_date),
        func.avg(FareObservation.total_fare), func.min(FareObservation.total_fare),
        func.max(FareObservation.total_fare),
    ).where(FareObservation.data_origin == ORIGIN))).one()
    return {
        "status": "READY", "records": total, "source": SOURCE, "data_origin": ORIGIN,
        "period_start": stats[0], "period_end": stats[1],
        "average_fare": round(float(stats[2]), 2), "min_fare": float(stats[3]), "max_fare": float(stats[4]),
        "note": "Historical reference only; excluded from live index and live fares.",
    }


@router.post("/backfill")
async def backfill_historical_snapshot(
    db: AsyncSession = Depends(get_db),
    _: dict = Depends(require_admin),
):
    """Import the local Kaggle workbook once. Subsequent calls are no-ops."""
    existing = await db.scalar(select(func.count()).select_from(FareObservation).where(
        FareObservation.data_origin == ORIGIN
    )) or 0
    if existing:
        return {"status": "ALREADY_IMPORTED", "records": existing, "source": SOURCE, "data_origin": ORIGIN}
    if not WORKBOOK.exists():
        raise HTTPException(status_code=404, detail=f"Historical workbook not found: {WORKBOOK.name}")

    try:
        workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(rows)]
        frame_rows = [dict(zip(headers, row)) for row in rows]
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"Could not read historical workbook: {exc}") from exc

    imported = 0
    now = datetime.now(timezone.utc)
    for row in frame_rows:
        origin_city, destination_city = str(row.get("Source", "")), str(row.get("Destination", ""))
        origin, destination = CITY_CODES.get(origin_city), CITY_CODES.get(destination_city)
        try:
            fare = float(row.get("Price"))
        except (TypeError, ValueError):
            fare = 0.0
        if not origin or not destination or fare <= 0:
            continue
        raw = f"{origin_city}|{destination_city}|{row.get('Date_of_Journey')}|{row.get('Airline')}|{fare}"
        db.add(FareObservation(
            collection_run_id="historical-kaggle-2019",
            origin=origin, destination=destination, route=f"{origin}-{destination}",
            airline=str(row.get("Airline", "Unknown"))[:10], flight_number=None,
            travel_date=_date(row.get("Date_of_Journey")), advance_days=0, collected_at=now,
            stops=0, cabin="ECONOMY", base_fare=float(fare), taxes=0.0, fees=0.0,
            total_fare=float(fare), currency="INR", availability_status="HISTORICAL",
            source=SOURCE, source_url="https://www.kaggle.com/datasets/nikhilmittal/flight-fare-prediction-mh/",
            data_origin=ORIGIN, collector_version="1.0",
            raw_hash=hashlib.sha256(raw.encode()).hexdigest(), quality_flags='["HISTORICAL_REFERENCE"]',
            is_valid=True,
        ))
        imported += 1
    await db.commit()
    return {"status": "IMPORTED", "records": imported, "source": SOURCE, "data_origin": ORIGIN}
